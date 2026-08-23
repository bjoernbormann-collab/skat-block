"""Modul 1 -- Import und Export als XLSX und CSV mit konfigurierbarem Mapping.

The mapping is column-header -> internal field, supplied per import run, so a
supplier list can be loaded without editing it first. FIELD_SPECS below is the
single definition of what a field is called, how it is parsed and how it is
written back out -- import, export and the mapping UI all read from it.

Import is transactional per run and defaults to a dry run: the caller gets the
full report first and only then decides to commit. Silently importing 4000 rows
of a mis-mapped spreadsheet is a very expensive mistake to undo.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.catalog import Article, Manufacturer, ProductGroup, Supplier
from app.models.enums import ArticleCondition, ArticleKind
from app.services.articles import (
    create_article,
    get_by_sku,
    set_oe_numbers,
    update_article,
)

# --------------------------------------------------------------------- parsing


class ParseError(ValueError):
    """Raised for a cell that cannot be turned into the target type."""


def parse_text(raw: str) -> str:
    return raw.strip()


def parse_decimal(raw: str) -> Decimal | None:
    """Accepts German and English notation: '1.234,56' and '1234.56'.

    Rejects anything else instead of guessing. A price that silently becomes 0
    is worse than a row that refuses to import.
    """
    text = raw.strip().replace(" ", "").replace(" ", "")
    if not text:
        return None
    text = text.replace("€", "").replace("EUR", "")
    if "," in text and "." in text:
        # Whichever separator comes last is the decimal one.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ParseError(f"„{raw}“ ist keine gültige Zahl.") from exc


def parse_quantity(raw: str) -> Decimal | None:
    value = parse_decimal(raw)
    return None if value is None else value.quantize(Decimal("0.001"))


def parse_money(raw: str) -> Decimal | None:
    value = parse_decimal(raw)
    return None if value is None else value.quantize(Decimal("0.0001"))


def parse_bool(raw: str) -> bool | None:
    text = raw.strip().lower()
    if text in ("", "-"):
        return None
    if text in ("1", "ja", "j", "true", "wahr", "x", "yes"):
        return True
    if text in ("0", "nein", "n", "false", "falsch", "no"):
        return False
    raise ParseError(f"„{raw}“ ist kein Ja/Nein-Wert.")


# Internal value -> German label. The export writes the German word (this is a
# German-language tool), and parse_condition accepts both spellings so an
# exported file still re-imports.
CONDITION_LABELS = {
    ArticleCondition.NEW.value: "neu",
    ArticleCondition.USED.value: "gebraucht",
    ArticleCondition.REFURBISHED.value: "generalüberholt",
}


def parse_condition(raw: str) -> str | None:
    text = raw.strip().lower()
    if not text:
        return None
    mapping = {
        "neu": ArticleCondition.NEW.value,
        "new": ArticleCondition.NEW.value,
        "gebraucht": ArticleCondition.USED.value,
        "used": ArticleCondition.USED.value,
        "generalüberholt": ArticleCondition.REFURBISHED.value,
        "generalueberholt": ArticleCondition.REFURBISHED.value,
        "refurbished": ArticleCondition.REFURBISHED.value,
    }
    if text not in mapping:
        raise ParseError(
            f"„{raw}“ ist kein bekannter Zustand (neu / gebraucht / generalüberholt)."
        )
    return mapping[text]


def format_value(value: Any) -> str:
    """Export side: German decimal notation, empty string for None."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value.normalize(), "f").replace(".", ",")
    if isinstance(value, bool):
        return "ja" if value else "nein"
    return str(value)


@dataclass(frozen=True)
class FieldSpec:
    """One importable/exportable article field."""

    key: str                       # internal attribute name
    label: str                     # German column header used on export
    parse: Callable[[str], Any]
    aliases: tuple[str, ...] = ()  # headers auto-recognised on import
    required: bool = False


FIELD_SPECS: tuple[FieldSpec, ...] = (
    FieldSpec("sku", "Artikelnummer", parse_text, ("sku", "artikelnr", "artikelnummer")),
    FieldSpec("title", "Titel", parse_text, ("titel", "title", "bezeichnung"),
              required=True),
    FieldSpec("ean", "EAN", parse_text, ("ean", "gtin", "barcode")),
    FieldSpec("mpn", "Herstellerteilenummer", parse_text,
              ("mpn", "hersteller-nr", "herstellernummer", "herstellerteilenummer")),
    FieldSpec("manufacturer_name", "Hersteller", parse_text,
              ("hersteller", "manufacturer", "marke")),
    FieldSpec("short_description", "Kurzbeschreibung", parse_text,
              ("kurzbeschreibung", "kurztext")),
    FieldSpec("long_description", "Langbeschreibung", parse_text,
              ("langbeschreibung", "beschreibung", "langtext")),
    FieldSpec("purchase_price_net", "EK netto", parse_money,
              ("ek", "ek netto", "einkaufspreis", "einkaufspreis netto")),
    FieldSpec("sales_price_gross", "VK brutto", parse_money,
              ("vk", "vk brutto", "verkaufspreis", "preis")),
    FieldSpec("vat_rate", "USt-Satz", parse_money, ("ust", "mwst", "steuersatz")),
    FieldSpec("weight_gram", "Gewicht (g)", parse_quantity, ("gewicht", "gewicht g")),
    FieldSpec("length_mm", "Länge (mm)", parse_quantity, ("länge", "laenge", "l")),
    FieldSpec("width_mm", "Breite (mm)", parse_quantity, ("breite", "b")),
    FieldSpec("height_mm", "Höhe (mm)", parse_quantity, ("höhe", "hoehe", "h")),
    FieldSpec("packaging_unit", "Verpackungseinheit", parse_quantity, ("ve", "vpe")),
    FieldSpec("hs_code", "Zolltarifnummer", parse_text, ("hs-code", "zolltarifnummer")),
    FieldSpec("country_of_origin", "Ursprungsland", parse_text,
              ("ursprungsland", "herkunftsland")),
    FieldSpec("supplier_name", "Lieferant", parse_text, ("lieferant", "supplier")),
    FieldSpec("supplier_sku", "Lieferanten-Artikelnr.", parse_text,
              ("lieferantennummer", "lief-artikelnr")),
    FieldSpec("product_group_code", "Warengruppe", parse_text,
              ("warengruppe", "gruppe", "kategorie")),
    FieldSpec("condition", "Zustand", parse_condition, ("zustand", "condition")),
    FieldSpec("min_stock", "Mindestbestand", parse_quantity,
              ("mindestbestand", "meldebestand")),
    FieldSpec("is_active", "Aktiv", parse_bool, ("aktiv", "active")),
    FieldSpec("oe_numbers", "OE-Nummern", parse_text,
              ("oe", "oe-nummern", "oe nummern", "originalnummer")),
    FieldSpec("internal_note", "Interne Notiz", parse_text, ("notiz", "bemerkung")),
)

SPECS_BY_KEY = {spec.key: spec for spec in FIELD_SPECS}

# Fields resolved to a foreign key rather than set directly.
RELATION_FIELDS = {"manufacturer_name", "supplier_name", "product_group_code"}
# Handled separately after the article exists.
DEFERRED_FIELDS = {"oe_numbers"}


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Guess header -> field from the aliases. The UI shows this as a starting
    point; the user can override every column."""
    mapping: dict[str, str] = {}
    for header in headers:
        needle = header.strip().lower()
        if not needle:
            continue
        for spec in FIELD_SPECS:
            if needle == spec.label.lower() or needle in spec.aliases:
                mapping[header] = spec.key
                break
    return mapping


# ---------------------------------------------------------------------- report


@dataclass
class RowResult:
    row_number: int
    action: str                       # "created" | "updated" | "skipped" | "error"
    sku: str = ""
    message: str = ""
    changed_fields: int = 0


@dataclass
class ImportReport:
    dry_run: bool
    rows: list[RowResult] = field(default_factory=list)

    @property
    def created(self) -> int:
        return sum(1 for row in self.rows if row.action == "created")

    @property
    def updated(self) -> int:
        return sum(1 for row in self.rows if row.action == "updated")

    @property
    def skipped(self) -> int:
        return sum(1 for row in self.rows if row.action == "skipped")

    @property
    def errors(self) -> list[RowResult]:
        return [row for row in self.rows if row.action == "error"]

    @property
    def ok(self) -> bool:
        return not self.errors


# ---------------------------------------------------------------------- reading


def read_table(data: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Read XLSX or CSV into (headers, rows-of-strings).

    Everything is read as text and converted by the field parsers. Letting
    openpyxl hand back floats for prices would reintroduce exactly the float
    problem the storage layer avoids.
    """
    lowered = filename.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    return _read_csv(data)


def _read_xlsx(data: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise ParseError(
            "Für XLSX-Dateien fehlt openpyxl. Installation: pip install -e ."
        ) from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []

    headers = [("" if cell is None else str(cell)).strip() for cell in header_row]
    rows: list[list[str]] = []
    for raw_row in rows_iter:
        cells = [_cell_to_text(cell) for cell in raw_row]
        if any(cell for cell in cells):
            rows.append(cells)
    workbook.close()
    return headers, rows


def _cell_to_text(cell: Any) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float):
        # openpyxl gives 4.99 as 4.99 but 4.0 as 4.0; keep it readable and let
        # parse_decimal do the real work.
        return repr(cell) if cell != int(cell) else str(int(cell))
    return str(cell).strip()


def _read_csv(data: bytes) -> tuple[list[str], list[list[str]]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes anything
        raise ParseError("Die Datei konnte nicht gelesen werden (Zeichensatz).")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        # German exports are semicolon-separated far more often than not.
        dialect = csv.excel
        dialect.delimiter = ";"

    reader = csv.reader(io.StringIO(text), dialect)
    try:
        headers = [cell.strip() for cell in next(reader)]
    except StopIteration:
        return [], []
    rows = [[cell.strip() for cell in row] for row in reader if any(c.strip() for c in row)]
    return headers, rows


# ---------------------------------------------------------------------- import


def _resolve_manufacturer(db: Session, name: str) -> int | None:
    if not name:
        return None
    row = db.scalar(select(Manufacturer).where(Manufacturer.name == name))
    if row is None:
        row = Manufacturer(name=name)
        db.add(row)
        db.flush()
    return row.id


def _resolve_supplier(db: Session, name: str) -> int | None:
    if not name:
        return None
    row = db.scalar(select(Supplier).where(Supplier.name == name))
    if row is None:
        row = Supplier(name=name)
        db.add(row)
        db.flush()
    return row.id


def _resolve_product_group(db: Session, code: str) -> int | None:
    """Unlike manufacturers and suppliers, groups are NOT created on the fly.

    A product group carries the eBay category mapping and the listing template.
    Auto-creating one from a typo in a spreadsheet would produce articles that
    look fine and cannot be published.
    """
    if not code:
        return None
    row = db.scalar(select(ProductGroup).where(ProductGroup.code == code))
    if row is None:
        raise ParseError(
            f"Warengruppe „{code}“ existiert nicht. Bitte zuerst anlegen."
        )
    return row.id


def import_articles(
    db: Session,
    *,
    headers: list[str],
    rows: list[list[str]],
    mapping: dict[str, str],
    dry_run: bool = True,
    update_existing: bool = True,
    user_id: int | None = None,
) -> ImportReport:
    """Import rows under the given header -> field mapping.

    With ``dry_run=True`` (the default) everything is rolled back to a savepoint
    at the end, so the report describes exactly what a real run would do.
    """
    report = ImportReport(dry_run=dry_run)

    mapped_keys = set(mapping.values())
    missing_required = [
        spec.label for spec in FIELD_SPECS
        if spec.required and spec.key not in mapped_keys
    ]
    if missing_required:
        raise ParseError(
            "Pflichtspalten fehlen in der Zuordnung: " + ", ".join(missing_required)
        )

    column_index = {header: index for index, header in enumerate(headers)}
    # Outer savepoint: the whole run is undone for a dry run, or when any row
    # failed. Importing "most of" a spreadsheet leaves nobody able to say what
    # actually went in.
    run_savepoint = db.begin_nested()

    for offset, row in enumerate(rows):
        # +2: one for the header line, one because humans count from 1.
        row_number = offset + 2
        # Per-row savepoint. Without it, a row that trips a database constraint
        # leaves the session unusable and every later row fails too, which turns
        # one bad cell into a report full of noise.
        row_savepoint = db.begin_nested()
        values: dict[str, Any] = {}
        try:
            for header, field_key in mapping.items():
                index = column_index.get(header)
                if index is None or index >= len(row):
                    continue
                spec = SPECS_BY_KEY.get(field_key)
                if spec is None:
                    continue
                parsed = spec.parse(row[index])
                if parsed is not None and parsed != "":
                    values[field_key] = parsed

            title = values.get("title", "")
            sku = values.get("sku", "")
            if not title and not sku:
                report.rows.append(RowResult(row_number, "skipped",
                                             message="Zeile ohne Titel und Artikelnummer"))
                continue

            oe_raw = values.pop("oe_numbers", "")
            values["manufacturer_id"] = _resolve_manufacturer(
                db, values.pop("manufacturer_name", ""))
            values["supplier_id"] = _resolve_supplier(db, values.pop("supplier_name", ""))
            values["product_group_id"] = _resolve_product_group(
                db, values.pop("product_group_code", ""))
            values = {key: value for key, value in values.items() if value is not None}

            existing = get_by_sku(db, sku) if sku else None
            if existing is not None:
                if not update_existing:
                    report.rows.append(RowResult(row_number, "skipped", sku=sku,
                                                 message="Artikel existiert bereits"))
                    continue
                changed = update_article(db, existing, values, user_id=user_id,
                                         source="import")
                article = existing
                report.rows.append(RowResult(row_number, "updated", sku=sku,
                                             changed_fields=changed))
            else:
                values.setdefault("kind", ArticleKind.STANDALONE.value)
                article = create_article(db, values, user_id=user_id, source="import")
                report.rows.append(RowResult(row_number, "created", sku=article.sku))

            if oe_raw:
                entries = [(part.strip(), "") for part in _split_list(oe_raw)]
                set_oe_numbers(db, article, entries)

            row_savepoint.commit()   # releases the savepoint, does not commit the run

        except Exception as exc:
            # Catching broadly is deliberate: a parse error, a constraint
            # violation and a missing product group should all become one
            # reportable bad row rather than aborting the run.
            row_savepoint.rollback()
            report.rows.append(RowResult(row_number, "error",
                                         sku=str(values.get("sku", "")),
                                         message=str(exc)))

    if dry_run or not report.ok:
        run_savepoint.rollback()
    else:
        run_savepoint.commit()
    return report


def _split_list(raw: str) -> list[str]:
    """OE numbers arrive comma-, semicolon- or pipe-separated depending on who
    exported the sheet."""
    for separator in (";", "|", ","):
        if separator in raw:
            return [part for part in raw.split(separator) if part.strip()]
    return [raw] if raw.strip() else []


# ---------------------------------------------------------------------- export


def export_rows(db: Session, articles: list[Article]) -> tuple[list[str], list[list[str]]]:
    """Build the export table. Round-trips: the export headers are exactly the
    labels the importer recognises, so an exported file can be edited and fed
    straight back in."""
    headers = [spec.label for spec in FIELD_SPECS]
    rows: list[list[str]] = []
    for article in articles:
        row: list[str] = []
        for spec in FIELD_SPECS:
            if spec.key == "manufacturer_name":
                row.append(article.manufacturer.name if article.manufacturer else "")
            elif spec.key == "supplier_name":
                row.append(article.supplier.name if article.supplier else "")
            elif spec.key == "product_group_code":
                row.append(article.product_group.code if article.product_group else "")
            elif spec.key == "condition":
                row.append(CONDITION_LABELS.get(article.condition, article.condition))
            elif spec.key == "oe_numbers":
                row.append("; ".join(
                    link.oe_number.number for link in article.oe_numbers
                ))
            else:
                row.append(format_value(getattr(article, spec.key, None)))
        rows.append(row)
    return headers, rows


def export_csv(db: Session, articles: list[Article]) -> bytes:
    headers, rows = export_rows(db, articles)
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL,
                        lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # utf-8-sig: without the BOM, Excel on macOS mangles every umlaut.
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(db: Session, articles: list[Article]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - depends on optional extra
        raise ParseError(
            "Für den XLSX-Export fehlt openpyxl. Installation: pip install -e ."
        ) from exc

    headers, rows = export_rows(db, articles)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Artikel"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"

    for index, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 12)
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = \
            min(width, 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
